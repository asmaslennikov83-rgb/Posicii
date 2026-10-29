from datetime import datetime, timedelta
from pathlib import Path
import csv
import json
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .config import MSK
from .wb import WBError


def period(now=None):
    now = now or datetime.now(MSK)
    end = now.astimezone(MSK).date() - timedelta(days=1)
    return end - timedelta(days=29), end


def barcodes(card):
    return {str(sku).strip() for size in card.get('sizes', [])
            for sku in size.get('skus', []) if str(sku).strip()}


def resolve(catalogs, query):
    """Expand exact barcode/vendor matches across BOTH catalogues until stable.

    A vendor's sizes and cards can have different nmIDs in each account.
    No prefix/substring match or numeric conversion of barcode strings.
    """
    vendors, skus = {query.strip().casefold()}, {query.strip()}
    matched = {key: {} for key in catalogs}
    changed = True
    while changed:
        changed = False
        for key, cards in catalogs.items():
            for card in cards:
                nm_id = int(card['nmID'])
                if nm_id in matched[key]:
                    continue
                vendor = str(card.get('vendorCode', '')).strip().casefold()
                codes = barcodes(card)
                if (vendor and vendor in vendors) or codes & skus:
                    matched[key][nm_id] = card
                    if vendor:
                        vendors.add(vendor)
                    skus.update(codes)
                    changed = True
    return {key: list(cards.values()) for key, cards in matched.items()}


def metric(row, key):
    value = row.get(key)
    if isinstance(value, dict):
        value = value.get('current')
    if value is None:
        return None
    if (isinstance(value, bool) or not isinstance(value, (int, float))
            or not math.isfinite(value) or value < 0):
        raise WBError('Изменился формат показателя ' + key)
    return value


def top(rows, key):
    # An absent frequency is not a real zero and must not enter a ranked list.
    unique = {}
    for row in rows:
        if metric(row, key) is None:
            continue
        text = row['text'].strip().casefold()
        if text in unique and unique[text] != row:
            raise WBError('WB вернул противоречивые дубли поискового запроса')
        unique[text] = row
    return sorted(unique.values(), key=lambda x: (-metric(x, key), x['text'].casefold()))[:10]


def read_complete_queries(folder, nm_id, start, end):
    """Optional complete data exported from an independently verified source.

    A manifest explicitly certifies coverage; a CSV alone is never assumed full.
    Files are local operator inputs, never fetched from arbitrary URLs.
    """
    manifest = folder / 'queries.json'
    if not manifest.exists():
        return None
    try:
        meta = json.loads(manifest.read_text(encoding='utf-8-sig'))
        if (meta.get('start') != start.isoformat() or meta.get('end') != end.isoformat()
                or meta.get('complete') is not True or not meta.get('source')):
            return None
        if nm_id not in meta.get('nm_ids', []):
            return None
        with (folder / 'queries.csv').open(encoding='utf-8-sig', newline='') as f:
            reader = csv.DictReader(f)
            required = {'nm_id', 'query', 'open_card', 'frequency', 'avg_position'}
            if not required <= set(reader.fieldnames or []):
                raise ValueError()
            rows = []
            for row in reader:
                if int(row['nm_id']) != nm_id:
                    continue
                if not row['query'].strip():
                    raise ValueError()
                item = {'nmId': nm_id, 'text': row['query'],
                             'openCard': float(row['open_card']), 'frequency': float(row['frequency']),
                             'avgPosition': float(row['avg_position']) if row['avg_position'] else None}
                for key in ('openCard', 'frequency', 'avgPosition'):
                    metric(item, key)
                rows.append(item)
        return rows, str(meta['source'])
    except (ValueError, OSError, KeyError, TypeError):
        raise WBError('Ошибка дополнительного источника запросов: проверьте queries.csv и queries.json') from None


def read_positions(folder):
    file = folder / 'positions.csv'
    if not file.exists():
        return {}
    result = {}
    try:
        with file.open(encoding='utf-8-sig', newline='') as f:
            reader = csv.DictReader(f)
            required = {'nm_id', 'query', 'organic_position', 'ad_position', 'measured_at', 'region', 'source'}
            if not required <= set(reader.fieldnames or []):
                raise ValueError()
            for row in reader:
                dt = datetime.fromisoformat(row['measured_at'])
                if dt.tzinfo is None or not row['source'].strip() or not row['region'].strip():
                    raise ValueError()
                for key in ('organic_position', 'ad_position'):
                    row[key] = float(row[key]) if row[key] else None
                    if row[key] is not None and (not math.isfinite(row[key]) or row[key] <= 0):
                        raise ValueError()
                key = (int(row['nm_id']), row['query'].strip().casefold())
                if key not in result or dt > datetime.fromisoformat(result[key]['measured_at']):
                    result[key] = row
        return result
    except (ValueError, OSError, KeyError, TypeError):
        raise WBError('Ошибка файла раздельных позиций: проверьте positions.csv') from None


HEADERS = ['Место', 'Артикул WB', 'Артикул продавца', 'Все ШК карточки', 'Запрос',
           'Переходы за 30 дней', 'Частотность за период WB', 'Средняя позиция WB за 30 дней',
           'Органическая позиция (замер)', 'Рекламная позиция (замер)', 'Дата замера позиций',
           'Регион замера', 'Источник позиций', 'Начало периода', 'Конец периода',
           'Охват запросов', 'Источник запросов', 'Примечание']


def text_cell(cell, value):
    cell.value = value
    # Store all external strings as literal text, including leading '=' in SKUs/queries.
    if isinstance(value, str):
        cell.data_type = 's'


def build_report(account, query, cards, api, data_dir: Path, start, end, limit=30):
    folder = data_dir / 'imports' / account.key
    warnings, groups, raw = [], [], []
    try:
        positions = read_positions(folder)
    except WBError as error:
        positions = {}
        warnings.append(str(error))
    for card in cards:
        nm_id = int(card['nmID'])
        try:
            complete = read_complete_queries(folder, nm_id, start, end)
            if complete is not None:
                rows, source = complete
                coverage = 'Полная выборка дополнительного источника'
            else:
                rows = api.queries(nm_id, start, end, limit)
                source = 'WB API: search-report/product/search-texts'
                coverage = f'До {limit} запросов WB, отобранных по переходам; полный охват не подтверждён'
                warnings.append(f'{nm_id}: топ частотности рассчитан по доступной выборке, не по всем запросам')
            raw.extend(rows)
            groups.append((card, top(rows, 'openCard'), top(rows, 'frequency'), coverage, source, ''))
            if positions and any((nm_id, row['text'].strip().casefold()) not in positions
                                 for row in top(rows, 'openCard') + top(rows, 'frequency')):
                warnings.append(f'{nm_id}: для части запросов нет отдельного замера позиций')
            if not rows:
                warnings.append(f'{nm_id}: WB не вернул запросы за период')
        except WBError as error:
            warnings.append(f'{nm_id}: {error}')
            groups.append((card, [], [], 'Нет данных', '', str(error)))
    if not cards:
        warnings.append('Карточка не найдена. Проверьте сопоставление ШК и доступ к каталогу')
    if not positions:
        warnings.append('Раздельные органические и рекламные позиции не подключены; средняя позиция WB их не заменяет')
    book = Workbook()
    book.remove(book.active)
    for index, title in ((1, 'ТОП-10 переходы'), (2, 'ТОП-10 частотность')):
        ws = book.create_sheet(title)
        notes = [f'{account.name} | Поиск: {query} | {start:%d.%m.%Y}–{end:%d.%m.%Y}',
                 'Частотность: по доступной выборке, если полный источник не подключён. ШК не суммируются.',
                 'Средняя позиция WB не является органической. Отдельные позиции — замер с указанными датой и регионом.',
                 'Статус: ' + ('; '.join(dict.fromkeys(warnings)) or 'Данные получены')]
        for note in notes:
            ws.append([None])
            text_cell(ws.cell(ws.max_row, 1), note)
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=len(HEADERS))
            ws.row_dimensions[ws.max_row].height = 42
        ws.append(HEADERS)
        for card, clicks, frequency, coverage, source, error in groups:
            selected = clicks if index == 1 else frequency
            if not selected:
                selected = [None]
            for rank, row in enumerate(selected, 1):
                nm_id = int(card['nmID'])
                pos = positions.get((nm_id, row['text'].strip().casefold()), {}) if row else {}
                values = [rank if row else None, str(nm_id), str(card.get('vendorCode', '')),
                          '\n'.join(sorted(barcodes(card))), row['text'] if row else '',
                          metric(row, 'openCard') if row else None,
                          metric(row, 'frequency') if row else None,
                          metric(row, 'avgPosition') if row else None,
                          pos.get('organic_position'), pos.get('ad_position'), pos.get('measured_at', ''),
                          pos.get('region', ''), pos.get('source', 'Не подключён'), start.isoformat(), end.isoformat(),
                          coverage, source, error or ('Нет запросов с этим показателем' if not row else '')]
                ws.append([None] * len(values))
                for col, value in enumerate(values, 1):
                    text_cell(ws.cell(ws.max_row, col), value)
        ws.freeze_panes = 'F6'
        ws.auto_filter.ref = f'A5:R{max(5, ws.max_row)}'
        for cell in ws[5]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[5].height = 48
        for col in 'ABCDEFGHIJKLMNOPQR':
            ws.column_dimensions[col].width = 22
        for col in ('D', 'E', 'P', 'Q', 'R'):
            ws.column_dimensions[col].width = 40
        for cells in ws.iter_rows(min_row=6):
            for cell in cells:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    stamp = datetime.now(MSK).strftime('%Y%m%d_%H%M%S_%f')
    output = data_dir / 'reports' / account.key
    output.mkdir(parents=True, exist_ok=True)
    file = output / f'cabinet_{account.key}_{stamp}.xlsx'
    book.save(file)
    # Raw responses contain no request headers or API tokens.
    file.with_suffix('.json').write_text(json.dumps({'query': query, 'start': start.isoformat(),
        'end': end.isoformat(), 'warnings': warnings, 'queries': raw}, ensure_ascii=False, indent=2), encoding='utf-8')
    return file, list(dict.fromkeys(warnings))
